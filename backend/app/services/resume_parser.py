import re
import os
from typing import Dict, Any, List
from pypdf import PdfReader
import docx

class ResumeParser:
    EMAIL_REGEX = r'[\w\.-]+@[\w\.-]+\.\w+'
    PHONE_REGEX = r'(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}'
    URL_REGEX = r'(?:https?://)?(?:www\.)?(?:linkedin\.com/in/[\w-]+|github\.com/[\w-]+)'

    COMMON_TECH_SKILLS = [
        "Python", "Java", "C++", "C", "C#", "JavaScript", "TypeScript", "Go", "Rust", "SQL",
        "HTML", "CSS", "React", "Next.js", "Vue", "Node.js", "Express", "FastAPI", "Django", "Flask",
        "Machine Learning", "Deep Learning", "Artificial Intelligence", "NLP", "Natural Language Processing",
        "Computer Vision", "LLMs", "Transformers", "PyTorch", "TensorFlow", "Scikit-Learn", "Keras",
        "Pandas", "NumPy", "OpenCV", "Hugging Face", "LangChain", "LlamaIndex",
        "Docker", "Kubernetes", "AWS", "Azure", "GCP", "Git", "GitHub", "Linux", "CI/CD",
        "PostgreSQL", "MySQL", "MongoDB", "Redis", "Elasticsearch", "REST API", "GraphQL",
        "Data Structures", "Algorithms", "System Design", "Microservices", "Agile", "Scrum"
    ]

    @classmethod
    def extract_text_from_file(cls, file_path: str, file_type: str) -> str:
        text = ""
        file_ext = os.path.splitext(file_path)[1].lower()

        if file_ext == ".pdf" or file_type == "pdf":
            try:
                reader = PdfReader(file_path)
                pages_text = []
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        pages_text.append(page_text)
                text = "\n".join(pages_text)
            except Exception as e:
                text = f"Error reading PDF: {str(e)}"

        elif file_ext == ".docx" or file_type == "docx":
            try:
                doc = docx.Document(file_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text]
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text and cell.text not in paragraphs:
                                paragraphs.append(cell.text)
                text = "\n".join(paragraphs)
            except Exception as e:
                text = f"Error reading DOCX: {str(e)}"

        elif file_ext == ".txt" or file_type == "txt":
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
            except Exception as e:
                text = f"Error reading TXT: {str(e)}"

        return text.strip()

    @classmethod
    def parse_resume(cls, raw_text: str) -> Dict[str, Any]:
        # Extract emails
        emails = re.findall(cls.EMAIL_REGEX, raw_text)
        email = emails[0] if emails else None

        # Extract phones
        phones = re.findall(cls.PHONE_REGEX, raw_text)
        phone = phones[0] if phones else None

        # Extract links
        links = re.findall(cls.URL_REGEX, raw_text, re.IGNORECASE)

        # Detect skills
        detected_skills = []
        lower_text = raw_text.lower()
        for skill in cls.COMMON_TECH_SKILLS:
            # Match whole words or boundary
            pattern = r'(?<![a-zA-Z0-9])' + re.escape(skill.lower()) + r'(?![a-zA-Z0-9])'
            if re.search(pattern, lower_text):
                detected_skills.append(skill)

        # Segment sections
        sections = cls._segment_sections(raw_text)

        return {
            "contact": {
                "email": email,
                "phone": phone,
                "links": links
            },
            "detected_skills": list(set(detected_skills)),
            "sections": sections
        }

    @staticmethod
    def _segment_sections(text: str) -> Dict[str, str]:
        section_headers = {
            "education": r'(?:education|academic\s+background|qualifications)',
            "experience": r'(?:experience|work\s+experience|employment|professional\s+experience|internships)',
            "projects": r'(?:projects|academic\s+projects|personal\s+projects)',
            "skills": r'(?:skills|technical\s+skills|core\s+competencies|technologies)',
            "certifications": r'(?:certifications|courses|achievements|honors|awards)'
        }

        results = {
            "education": "",
            "experience": "",
            "projects": "",
            "skills": "",
            "certifications": ""
        }

        # Simple line by line section splitter
        lines = text.split("\n")
        current_section = None
        current_lines = []

        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue

            matched_sec = None
            if len(stripped) < 40:  # Header likely
                for sec, pattern in section_headers.items():
                    if re.match(r'^[\s#*_-]*' + pattern + r'[\s:*_-]*$', stripped, re.IGNORECASE):
                        matched_sec = sec
                        break

            if matched_sec:
                if current_section:
                    results[current_section] = "\n".join(current_lines).strip()
                current_section = matched_sec
                current_lines = []
            else:
                if current_section:
                    current_lines.append(stripped)

        if current_section:
            results[current_section] = "\n".join(current_lines).strip()

        return results
